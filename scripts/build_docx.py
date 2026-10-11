#!/usr/bin/env python3
"""Build a journal club outline .docx from a JSON spec.

  build_docx.py outline.json

Spec format:
{
  "output": "/abs/path/Outline.docx",
  "header": {"title": "...", "authors": "...", "institution": "...", "journal": "..."},
  "blocks": [
    {"h1": "Results", "time": "~5 min"},            # section heading (time is optional)
    {"h2": "HG recurrence-free survival"},           # sub-heading
    {"bullets": ["text", {"text": "parent", "sub": ["child", "child"]}]},
    {"figure": "/abs/path/fig1.png", "caption": "Figure 1 - ...", "width_in": 5.0},
    {"callout": "One-line take-home message"}        # shaded box
  ]
}
Inline **bold** is supported in any text field.
"""
import json
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x66, 0x66, 0x66)


def add_runs(par, text, size=None, color=None, italic=False):
    for i, chunk in enumerate(re.split(r"\*\*", text)):
        if not chunk:
            continue
        r = par.add_run(chunk)
        r.bold = i % 2 == 1
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color


def tight(par, before=0, after=2):
    par.paragraph_format.space_before = Pt(before)
    par.paragraph_format.space_after = Pt(after)


def shade(par, hex_fill):
    ppr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    ppr.append(shd)


def bullets(doc, items, level=1):
    style = "List Bullet" if level == 1 else f"List Bullet {min(level, 3)}"
    for it in items:
        text, sub = (it, None) if isinstance(it, str) else (it["text"], it.get("sub"))
        p = doc.add_paragraph(style=style)
        tight(p, 0, 1)
        p.paragraph_format.left_indent = Inches(0.25 + 0.3 * (level - 1))
        p.paragraph_format.first_line_indent = Inches(-0.2)
        add_runs(p, text)
        if sub:
            bullets(doc, sub, level + 1)


def main(spec_path):
    spec = json.load(open(spec_path))
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.75)
        s.top_margin = s.bottom_margin = Inches(0.6)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for name, size in (("Heading 1", 14), ("Heading 2", 12)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = ACCENT

    h = spec["header"]
    p = doc.add_paragraph()
    tight(p, 0, 4)
    add_runs(p, h["title"], size=15, color=ACCENT)
    p.runs[0].bold = True
    for label, key in (("Lead author", "authors"), ("Institution", "institution"), ("Journal", "journal")):
        if h.get(key):
            p = doc.add_paragraph()
            tight(p, 0, 0)
            add_runs(p, f"**{label}:** {h[key]}")

    for b in spec["blocks"]:
        if "h1" in b:
            p = doc.add_heading(level=1)
            tight(p, 10, 2)
            add_runs(p, b["h1"])
            if b.get("time"):
                add_runs(p, f"   ({b['time']})", size=10, color=GREY, italic=True)
        elif "h2" in b:
            p = doc.add_heading(level=2)
            tight(p, 6, 1)
            add_runs(p, b["h2"])
        elif "bullets" in b:
            bullets(doc, b["bullets"])
        elif "figure" in b:
            p = doc.add_paragraph()
            tight(p, 4, 0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(b["figure"], width=Inches(b.get("width_in", 5.0)))
            if b.get("caption"):
                c = doc.add_paragraph()
                tight(c, 0, 6)
                c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(c, b["caption"], size=9, color=GREY, italic=True)
        elif "callout" in b:
            p = doc.add_paragraph()
            tight(p, 4, 4)
            shade(p, "E8EEF6")
            add_runs(p, b["callout"])
        else:
            sys.exit(f"Unknown block: {b}")

    doc.save(spec["output"])
    print(spec["output"])


if __name__ == "__main__":
    main(sys.argv[1])
